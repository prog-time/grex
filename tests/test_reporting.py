"""Tests for monitoring report generation and scheduled delivery.

The report-file path globals are patched to a tmp dir; DB-touching code runs
against a per-test temporary SQLite database (globals restored afterwards).
Channel sending is mocked — these tests never hit Telegram/SMTP.
"""
from __future__ import annotations

from datetime import datetime
from unittest.mock import AsyncMock, patch

import pytest_asyncio

from services import reporting


@pytest_asyncio.fixture
async def db(tmp_path):
    import models.alert_source  # noqa: F401 — register tables
    import models.parse_target  # noqa: F401
    import models.parsed_product  # noqa: F401
    import models.price_snapshot  # noqa: F401
    from models import job as job_mod

    saved_engine, saved_session = job_mod.engine, job_mod.async_session
    await job_mod.init_db(f"sqlite:///{tmp_path}/rep.db")
    try:
        yield job_mod.async_session
    finally:
        from models.job import close_db
        await close_db()
        job_mod.async_session = saved_session
        job_mod.engine = saved_engine


def test_frequency_seconds_matches_alert_frequencies():
    from admin.resources.alert_sources import _FREQUENCIES

    assert set(reporting.FREQUENCY_SECONDS) == set(_FREQUENCIES)


async def _generate(tmp_path):
    """Run generate_report() against report globals pointed at tmp_path."""
    latest = tmp_path / "monitoring_latest.xlsx"
    with patch.object(reporting, "_REPORTS_DIR", str(tmp_path)), \
            patch.object(reporting, "_LATEST", str(latest)):
        path = await reporting.generate_report()
    return path, latest


class TestGenerate:
    async def test_products_sheet_only_active_monitored(self, db, tmp_path):
        from models.parse_target import ParseTarget
        from models.parsed_product import ParsedProduct

        async with db() as s:
            s.add(ParseTarget(id="t1", url="https://ozon.ru/product/x-1", marketplace="ozon", monitor_enabled=True))
            s.add(ParseTarget(id="t2", url="https://ozon.ru/product/y-2", marketplace="ozon", monitor_enabled=False))
            s.add(ParsedProduct(url="https://ozon.ru/product/x-1", title="Товар X", sku="SKU1",
                                current_price=100.0, original_price=150.0, marketplace="ozon"))
            await s.commit()

        path, latest = await _generate(tmp_path)

        assert path == str(latest) and latest.exists()
        from openpyxl import load_workbook
        wb = load_workbook(latest)
        assert wb.sheetnames == ["Товары", "История цен", "Алерты", "Сводка"]
        ws = wb["Товары"]
        values = [[c.value for c in row] for row in ws.iter_rows()]
        assert values[2][0] == "ID товара" and values[2][1] == "Товар"  # header row (after title + blank)
        data_rows = [r for r in values[3:] if r[0]]
        assert len(data_rows) == 1                 # only the active monitored target
        assert data_rows[0][1] == "Товар X" and data_rows[0][2] == "SKU1"
        assert data_rows[0][5] == 100.0 and data_rows[0][6] == 150.0   # current / original price
        assert data_rows[0][8] == 33.33                                # discount %

    async def test_history_and_alert_sheets_and_period_window(self, db, tmp_path):
        from datetime import timedelta

        from models.parse_target import ParseTarget
        from models.parsed_product import ParsedProduct
        from models.price_snapshot import PriceAlert, PriceSnapshot

        url = "https://ozon.ru/product/x-1"
        now = datetime.utcnow()
        async with db() as s:
            s.add(ParseTarget(id="t1", url=url, marketplace="ozon", monitor_enabled=True))
            s.add(ParsedProduct(url=url, title="Товар X", current_price=90.0, original_price=150.0, marketplace="ozon"))
            # two in-window snapshots + one stale (outside the 30-day default window)
            s.add(PriceSnapshot(url=url, marketplace="ozon", current_price=120.0, parsed_at=now - timedelta(days=5)))
            s.add(PriceSnapshot(url=url, marketplace="ozon", current_price=90.0, parsed_at=now - timedelta(days=1)))
            s.add(PriceSnapshot(url=url, marketplace="ozon", current_price=999.0, parsed_at=now - timedelta(days=40)))
            s.add(PriceAlert(url=url, marketplace="ozon", title="Товар X", kind="drop",
                             old_price=120.0, new_price=90.0, change_pct=-25.0))
            await s.commit()

        _, latest = await _generate(tmp_path)
        from openpyxl import load_workbook
        wb = load_workbook(latest)

        hist = [[c.value for c in r] for r in wb["История цен"].iter_rows()]
        hist_rows = [r for r in hist[1:] if r[0]]
        assert len(hist_rows) == 2                 # stale 40-day snapshot excluded
        assert {r[3] for r in hist_rows} == {120.0, 90.0}

        alerts = [[c.value for c in r] for r in wb["Алерты"].iter_rows()]
        alert_rows = [r for r in alerts[1:] if r[0]]
        assert len(alert_rows) == 1 and alert_rows[0][3] == "Падение"

        # Products sheet stats are computed from the in-window snapshots only.
        prod = [[c.value for c in r] for r in wb["Товары"].iter_rows()][3]
        assert prod[9] == 90.0 and prod[10] == 120.0   # min / max price in window
        assert prod[15] == 2                            # observations (stale excluded)
        assert prod[16] == 1                            # alerts ↓


class TestEnsureFresh:
    async def test_generates_when_missing(self, tmp_path):
        latest = tmp_path / "none.xlsx"
        with patch.object(reporting, "_LATEST", str(latest)), \
                patch.object(reporting, "generate_report", new=AsyncMock(return_value=str(latest))) as gen:
            await reporting.ensure_report_fresh()
        gen.assert_awaited_once()

    async def test_skips_when_fresh(self, tmp_path):
        latest = tmp_path / "fresh.xlsx"
        latest.write_text("x")  # just created → age ~0 < interval
        with patch.object(reporting, "_LATEST", str(latest)), \
                patch.object(reporting, "generate_report", new=AsyncMock()) as gen:
            await reporting.ensure_report_fresh()
        gen.assert_not_awaited()


class TestDeliver:
    async def test_delivers_to_due_only_and_updates_last_sent(self, db, tmp_path):
        from sqlalchemy import select

        from models.alert_source import AlertSource

        report = tmp_path / "r.xlsx"
        report.write_text("x")
        now = datetime.utcnow()
        async with db() as s:
            s.add(AlertSource(id="a1", type="telegram", recipient="@u", chat_id="111",
                              frequency="hourly", last_sent_at=None))           # due
            s.add(AlertSource(id="a2", type="telegram", recipient="@v", chat_id="222",
                              frequency="hourly", last_sent_at=now))            # not due
            await s.commit()

        with patch.object(reporting, "_LATEST", str(report)), \
                patch.object(reporting, "_send_one", new=AsyncMock(return_value=True)) as send:
            sent = await reporting.deliver_due()

        assert sent == 1 and send.await_count == 1
        async with db() as s:
            by_id = {r.id: r.last_sent_at for r in (await s.execute(select(AlertSource))).scalars().all()}
        assert by_id["a1"] is not None        # due source marked as sent
        assert by_id["a2"] == now             # not-due source untouched

    async def test_does_not_update_last_sent_on_failure(self, db, tmp_path):
        from sqlalchemy import select

        from models.alert_source import AlertSource

        report = tmp_path / "r.xlsx"
        report.write_text("x")
        async with db() as s:
            s.add(AlertSource(id="a1", type="telegram", recipient="@u", chat_id="1", frequency="hourly", last_sent_at=None))
            await s.commit()

        with patch.object(reporting, "_LATEST", str(report)), \
                patch.object(reporting, "_send_one", new=AsyncMock(return_value=False)):
            sent = await reporting.deliver_due()

        assert sent == 0
        async with db() as s:
            src = (await s.execute(select(AlertSource))).scalars().first()
        assert src.last_sent_at is None       # not marked sent when delivery failed

    async def test_no_report_means_no_delivery(self, db, tmp_path):
        from models.alert_source import AlertSource

        async with db() as s:
            s.add(AlertSource(id="a1", type="telegram", recipient="@u", chat_id="1", frequency="hourly"))
            await s.commit()

        missing = tmp_path / "absent.xlsx"
        with patch.object(reporting, "_LATEST", str(missing)), \
                patch.object(reporting, "_send_one", new=AsyncMock(return_value=True)) as send:
            sent = await reporting.deliver_due()
        assert sent == 0 and send.await_count == 0
