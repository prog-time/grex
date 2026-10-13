"""Admin API — product link extraction routes."""
from __future__ import annotations

import io

from fastapi import Request
from fastapi.responses import JSONResponse, StreamingResponse
from loguru import logger


class LinksRouter:

    @staticmethod
    async def extract_links(request: Request) -> JSONResponse:
        try:
            from core.exceptions import UnsupportedMarketplaceError  # noqa: PLC0415
            from crawlers.registry import CrawlerRegistry  # noqa: PLC0415

            body = await request.json()
            url: str = (body.get("url") or "").strip()
            max_pages: int = min(int(body.get("max_pages") or 10), 50)

            if not url:
                return JSONResponse({"ok": False, "error": "URL обязателен"}, status_code=422)

            try:
                marketplace = CrawlerRegistry.detect_marketplace(url)
            except UnsupportedMarketplaceError:
                return JSONResponse(
                    {"ok": False, "error": f"Маркетплейс не распознан для URL: {url}"},
                    status_code=422,
                )

            crawler_cls = CrawlerRegistry.get(marketplace)
            crawler = crawler_cls()

            logger.info(f"Link extraction started: {marketplace} {url} max_pages={max_pages}")
            links = await crawler.get_product_links(url, max_pages=max_pages)
            logger.info(f"Link extraction done: {marketplace} → {len(links)} links")

            return JSONResponse({
                "ok": True,
                "marketplace": marketplace,
                "url": url,
                "total": len(links),
                "links": links,
            })

        except NotImplementedError as e:
            return JSONResponse({"ok": False, "error": str(e)}, status_code=400)
        except Exception as exc:  # noqa: BLE001
            logger.error(f"Link extraction error: {exc}")
            return JSONResponse({"ok": False, "error": str(exc)}, status_code=500)

    @staticmethod
    async def download_excel(request: Request):
        try:
            import openpyxl  # noqa: PLC0415
            from openpyxl.styles import Alignment, Font, PatternFill  # noqa: PLC0415

            body = await request.json()
            links: list[str] = body.get("links") or []
            marketplace: str = body.get("marketplace") or "links"

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Ссылки"

            # Header row
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            ws.append(["#", "URL"])
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            ws.column_dimensions["A"].width = 7
            ws.column_dimensions["B"].width = 90

            for i, link in enumerate(links, start=1):
                ws.append([i, link])

            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)

            filename = f"links_{marketplace}.xlsx"
            headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
            return StreamingResponse(
                buf,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers=headers,
            )

        except Exception as exc:  # noqa: BLE001
            logger.error(f"Excel export error: {exc}")
            return JSONResponse({"error": str(exc)}, status_code=500)

    @staticmethod
    async def download_csv(request: Request):
        try:
            import csv as _csv  # noqa: PLC0415

            body = await request.json()
            links: list[str] = body.get("links") or []
            marketplace: str = body.get("marketplace") or "links"

            buf = io.StringIO()
            writer = _csv.writer(buf)
            writer.writerow(["#", "URL"])
            for i, link in enumerate(links, start=1):
                writer.writerow([i, link])

            content = "\ufeff" + buf.getvalue()  # UTF-8 BOM for Excel compatibility
            filename = f"links_{marketplace}.csv"
            headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
            return StreamingResponse(
                iter([content.encode("utf-8")]),
                media_type="text/csv; charset=utf-8",
                headers=headers,
            )

        except Exception as exc:  # noqa: BLE001
            logger.error(f"CSV export error: {exc}")
            return JSONResponse({"error": str(exc)}, status_code=500)


def register_links_routes(app) -> None:  # type: ignore[no-untyped-def]
    from fastapi import Depends  # noqa: PLC0415

    from admin.resources.auth import require_admin  # noqa: PLC0415
    _auth = [Depends(require_admin)]
    router = LinksRouter()
    app.add_api_route(
        "/admin/links/extract",
        router.extract_links,
        methods=["POST"],
        tags=["admin"],
        dependencies=_auth,
    )
    app.add_api_route(
        "/admin/links/download-excel",
        router.download_excel,
        methods=["POST"],
        tags=["admin"],
        dependencies=_auth,
    )
    app.add_api_route(
        "/admin/links/download-csv",
        router.download_csv,
        methods=["POST"],
        tags=["admin"],
        dependencies=_auth,
    )
